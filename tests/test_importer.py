import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from academic_workload.importers.weekly_workbook import preview_workbook


class WorkbookPreviewTests(unittest.TestCase):
    def test_maps_actual_from_m_not_ai_and_skips_future_week_actual(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Weekly Workload"
        ws["D6"] = "Teaching (contact hours)"
        ws["A7"] = "Week No."
        ws["A8"] = 1
        ws["B8"] = datetime(2026, 9, 19)
        ws["C8"] = datetime(2026, 9, 24)
        ws["D8"] = 3
        ws["E8"] = 1
        ws["M8"] = 4
        ws["AI8"] = "=L8"  # Incorrectly labeled 'Actual' in workbook.
        ws["N8"] = 4
        ws["O8"] = 5
        ws["P8"] = 4
        ws["Q8"] = 2
        ws["R8"] = "Quality Assurance"
        ws["S8"] = 4
        ws["T8"] = 2
        ws["U8"] = 8
        ws["V8"] = 4
        ws["W8"] = 4
        ws["X8"] = 9
        ws["Y8"] = 4
        ws["Z8"] = 3
        ws["AA8"] = 8
        ws["AB8"] = 11
        ws["AH8"] = 10
        ws["A9"] = 2
        ws["B9"] = datetime(2026, 9, 26)
        ws["C9"] = datetime(2026, 10, 1)
        ws["D9"] = 3
        ws["E9"] = 1
        ws["AH9"] = 10
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "synthetic.xlsx"
            wb.save(path)
            result = preview_workbook(path, target_weekly_hours=40)
        self.assertEqual(len(result["weeks"]), 2)
        first, second = result["weeks"]
        self.assertEqual(first["actual_teaching_hours"], 4)
        self.assertEqual(first["planned_teaching_hours"], 4)
        self.assertEqual(first["actual_total_hours"], 40)
        self.assertEqual(first["teaching_gap"], 6)
        self.assertEqual(second["workload_status"], "Not Recorded")
        self.assertEqual(second["actual_recorded"], 0)


if __name__ == "__main__":
    unittest.main()
